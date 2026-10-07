# -*- coding: utf-8 -*-
"""
Metas de E-commerce para o ano de 2026.
Carregadas diretamente da pasta oficial do projeto: 'metas ecommerce'.
Permite atualização automática ao editar ou substituir a planilha Excel na pasta,
sem qualquer dependência de caminhos externos ou pastas fora do projeto.
"""
import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PASTA_METAS = os.path.join(BASE_DIR, "metas ecommerce")
EXCEL_PATH = os.path.join(PASTA_METAS, "Metas ecommerce 2026.xlsx")

_cache_mtime = None
_cache_metas = []
_cache_mapa = {}


def carregar_metas_excel(caminho=EXCEL_PATH):
    """
    Carrega as metas diretamente da planilha Excel dentro da pasta 'metas ecommerce'.
    Utiliza verificação de data de modificação (mtime) para recarregar automaticamente caso a planilha seja alterada.
    """
    global _cache_mtime, _cache_metas, _cache_mapa
    if not os.path.exists(caminho):
        return _cache_metas

    try:
        mtime = os.path.getmtime(caminho)
        if _cache_mtime == mtime and _cache_metas:
            return _cache_metas

        wb = openpyxl.load_workbook(caminho, data_only=True)
        sheet_name = 'METAS ECOMMERCE_2026' if 'METAS ECOMMERCE_2026' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        data = []
        for r in rows[1:]:
            if not r or not any(r[:6]):
                continue
            mes_nome = str(r[0]).strip() if r[0] else ''
            data_val = r[1]
            vendedor = str(r[2]).strip() if r[2] else ''
            vlr_meta = round(float(r[3]), 2) if r[3] is not None else 0.0
            dias_mes = int(r[4]) if r[4] is not None else 0
            vlr_dia = round(float(r[5]), 2) if r[5] is not None else 0.0

            mes_num = None
            if hasattr(data_val, 'month'):
                mes_num = data_val.month
            elif isinstance(data_val, (int, float)):
                mes_num = int(data_val)
            else:
                meses_map = {
                    'janeiro': 1, 'fevereiro': 2, 'março': 3, 'marco': 3,
                    'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7,
                    'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12
                }
                mes_num = meses_map.get(mes_nome.lower())

            if vendedor and vlr_meta > 0 and mes_num:
                data.append({
                    'ano': 2026,
                    'mes': mes_num,
                    'mes_nome': mes_nome,
                    'vendedor': vendedor,
                    'vlr_meta': vlr_meta,
                    'dias_mes': dias_mes,
                    'vlr_dia': vlr_dia
                })

        if data:
            _cache_metas = data
            _cache_mtime = mtime
            _cache_mapa = {(item["vendedor"].upper(), item["mes"]): item for item in _cache_metas}
    except Exception as e:
        print("[METAS] Aviso ao carregar Excel de metas:", e)

    return _cache_metas


# Carga inicial na importação do módulo
_METAS_INICIAIS = carregar_metas_excel()

# Constantes expostas para compatibilidade total com o sistema
METAS_ECOMMERCE_2026 = _METAS_INICIAIS
METAS_ECOMMERCE_2026_V2 = _METAS_INICIAIS


def _obter_dados_atualizados():
    metas = carregar_metas_excel()
    return metas if metas else _METAS_INICIAIS


def get_meta(vendedor: str, mes: int, versao: int = 1) -> dict:
    """
    Retorna o registro da meta de um vendedor/canal para o mês indicado (1 a 12).
    Retorna dict com chaves: ano, mes, mes_nome, vendedor, vlr_meta, dias_mes, vlr_dia.
    """
    carregar_metas_excel()
    return _cache_mapa.get((vendedor.strip().upper(), mes))


def get_meta_valor(vendedor: str, mes: int, versao: int = 1) -> float:
    """Retorna o valor financeiro mensal da meta (R$)."""
    m = get_meta(vendedor, mes, versao)
    return m["vlr_meta"] if m else 0.0


def get_meta_diaria(vendedor: str, mes: int, versao: int = 1) -> float:
    """Retorna a meta diária (R$/dia)."""
    m = get_meta(vendedor, mes, versao)
    return m["vlr_dia"] if m else 0.0


def get_total_meta_mes(mes: int, versao: int = 1, incluir_parlux: bool = False) -> float:
    """
    Retorna a soma de metas dos canais para o mês (1 a 12).
    Por diretriz da Diretoria de E-commerce, PARLUX não é considerada na meta geral/consolidada.
    """
    metas = _obter_dados_atualizados()
    if not incluir_parlux:
        return round(sum(item["vlr_meta"] for item in metas if item["mes"] == mes and item["vendedor"].upper() != "PARLUX"), 2)
    return round(sum(item["vlr_meta"] for item in metas if item["mes"] == mes), 2)


def get_metas_mes(mes: int, versao: int = 1) -> list:
    """Retorna lista com todos os canais para o mês especificado."""
    metas = _obter_dados_atualizados()
    return [item for item in metas if item["mes"] == mes]


def get_metas_vendedor(vendedor: str, versao: int = 1) -> list:
    """Retorna lista com os 12 meses de um determinado vendedor/canal."""
    metas = _obter_dados_atualizados()
    return [item for item in metas if item["vendedor"].upper() == vendedor.strip().upper()]


def get_resumo_anual(versao: int = 1) -> dict:
    """Retorna resumo anual agrupado por canal."""
    metas = _obter_dados_atualizados()
    resumo = {}
    for item in metas:
        v = item["vendedor"]
        if v not in resumo:
            resumo[v] = {"total_meta": 0.0, "meses": 0}
        resumo[v]["total_meta"] = round(resumo[v]["total_meta"] + item["vlr_meta"], 2)
        resumo[v]["meses"] += 1
    return resumo


if __name__ == "__main__":
    print("=== RESUMO ANUAL DE METAS 2026 (CARREGADAS DE 'metas ecommerce') ===")
    resumo = get_resumo_anual()
    total_geral = 0.0
    for canal, dados in resumo.items():
        total_geral += dados["total_meta"]
        print(f"  {canal:<15}: R$ {dados['total_meta']:>12,.2f}")
    print("-" * 35)
    print(f"  {'TOTAL GERAL':<15}: R$ {total_geral:>12,.2f}")
