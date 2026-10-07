# -*- coding: utf-8 -*-
"""
Script utilitário para exportação e consolidação das Metas de E-commerce.
Lê a planilha oficial diretamente da pasta local do projeto: 'metas ecommerce'.
"""
import openpyxl
import json
import os

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PASTA_METAS = os.path.join(BASE_DIR, "metas ecommerce")
excel_path = os.path.join(PASTA_METAS, "Metas ecommerce 2026.xlsx")

wb = openpyxl.load_workbook(excel_path, data_only=True)

def parse_sheet(sheet_name):
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for r in rows[1:]:
        if not r or not any(r[:6]):
            continue
        mes_nome = str(r[0]).strip() if r[0] else ''
        data_val = r[1]
        vendedor = str(r[2]).strip() if r[2] else ''
        vlr_meta = round(float(r[3]), 2) if r[3] is not None else 0.0
        dias_mes = int(r[4]) if r[4] is not None else 0
        vlr_dia = round(float(r[5]), 2) if r[5] is not None else 0.0
        
        mes_num = None
        if hasattr(data_val, 'month'):
            mes_num = data_val.month
        elif isinstance(data_val, (int, float)):
            mes_num = int(data_val)
        else:
            meses_map = {
                'janeiro': 1, 'fevereiro': 2, 'março': 3, 'marco': 3,
                'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7,
                'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12
            }
            mes_num = meses_map.get(mes_nome.lower())

        if vendedor and vlr_meta > 0 and mes_num:
            data.append({
                'ano': 2026,
                'mes': mes_num,
                'mes_nome': mes_nome,
                'vendedor': vendedor,
                'vlr_meta': vlr_meta,
                'dias_mes': dias_mes,
                'vlr_dia': vlr_dia
            })
    return data

d1 = parse_sheet('METAS ECOMMERCE_2026')
d2 = parse_sheet('METAS ECOMMERCE_2026 (2)')
if not d2:
    d2 = d1

# Salva JSON consolidado
json_path = os.path.join(BASE_DIR, 'metas_ecommerce_2026.json')
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump({'metas_ecommerce_2026': d1, 'metas_ecommerce_2026_v2': d2}, f, indent=2, ensure_ascii=False)

print(f"Metas exportadas com sucesso a partir de: {excel_path}")
print(f"Total registros na aba principal: {len(d1)}")
